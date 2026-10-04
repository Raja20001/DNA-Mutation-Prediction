"""
Generates a publication-grade academic Word document (.docx) from RESEARCH_PAPER.md
with IEEE/Academic styling: custom fonts, headings, tables, borders, callouts, and footers.
"""
from pathlib import Path
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding in twips (1/20th of a point)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    """Set subtle horizontal borders for academic tables."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="1B365D"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1B365D"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, line_spacing=1.15, bold_prefix=None):
    """Add a paragraph with parsed bold/italic markdown inline formatting."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    # Simple regex-based tokenization for **bold** and *italic*
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
            run = p.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
            run = p.add_run(tok[1:-1])
            run.italic = True
        elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
            run = p.add_run(tok[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
        else:
            p.add_run(tok)
    return p

def create_research_docx(md_path: Path, output_path: Path):
    doc = Document()

    # Page Margins (Normal 1 inch)
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        
        # Add Header & Footer
        header = sec.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("M.E. Thesis Research | DNA Variant Classical-Quantum Hybrid Framework")
        hr.font.name = 'Calibri'
        hr.font.size = Pt(8.5)
        hr.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run("Academic Research Dossier — Strictly for Research & Evaluation")
        fr.font.name = 'Calibri'
        fr.font.size = Pt(8.5)
        fr.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

    # Configure Default Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)

    # Read Markdown Lines
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        # Parse markdown table
        rows_data = []
        for tl in tbl_lines:
            tl = tl.strip()
            if not tl or not tl.startswith('|'):
                continue
            cells = [c.strip() for c in tl.split('|')[1:-1]]
            # Check if separator row
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            rows_data.append(cells)

        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        for row_idx, row in enumerate(rows_data):
            for col_idx in range(num_cols):
                cell_text = row[col_idx] if col_idx < len(row) else ""
                cell = table.cell(row_idx, col_idx)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=140, right=140)

                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05

                # Clean markdown formatting inside table
                cleaned_text = cell_text.replace('**', '').replace('*', '')
                run = p.add_run(cleaned_text)
                run.font.name = 'Calibri'
                run.font.size = Pt(9.5)

                if row_idx == 0:
                    set_cell_background(cell, "1B365D")
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "F8F9FA")
                    else:
                        set_cell_background(cell, "FFFFFF")
                    run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)

    def flush_code(c_lines):
        if not c_lines:
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F4F6F9")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        full_code = "\n".join(c_lines)
        run = p.add_run(full_code)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Code block handling
        if line.startswith('```'):
            if in_code_block:
                flush_code(code_lines)
                code_lines = []
                in_code_block = False
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table handling
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        stripped = line.strip()

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == '---':
            i += 1
            continue

        # Document Main Title (# Title without number)
        if line.startswith('# ') and not re.match(r'^#\s+\d+\.', line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(line[2:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Author block subtitle
            p_auth = doc.add_paragraph()
            p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_auth.paragraph_format.space_after = Pt(18)
            r_auth = p_auth.add_run("M.E. Computer Science Research Thesis Documentation\nAnna University / Academic Postgraduate Research Center")
            r_auth.font.name = 'Calibri'
            r_auth.font.size = Pt(11)
            r_auth.font.color.rgb = RGBColor(0x56, 0x65, 0x73)
            r_auth.italic = True
            i += 1
            continue

        # Heading 1 (## Section or # 1. Section)
        if line.startswith('## ') or (line.startswith('# ') and re.match(r'^#\s+\d+\.', line)):
            h_text = line[3:].strip() if line.startswith('## ') else line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            i += 1
            continue

        # Heading 2 (### Subsection)
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[4:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0x2C, 0x52, 0x82)
            i += 1
            continue

        # Heading 3 (#### Sub-subsection)
        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line[5:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
            i += 1
            continue

        # Blockquote (> text)
        if line.startswith('> '):
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "F0F4F8")
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            # Add left navy border to callout box
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/>'
                f'  <w:top w:val="none"/>'
                f'  <w:bottom w:val="none"/>'
                f'  <w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            quote_text = line[2:].strip()
            add_styled_paragraph(doc, quote_text, space_after=2)
            # Transfer text from added paragraph to cell
            last_p = doc.paragraphs[-1]
            p.clear()
            for r in last_p.runs:
                nr = p.add_run(r.text)
                nr.bold = r.bold
                nr.italic = True
                nr.font.name = 'Calibri'
                nr.font.size = Pt(10.5)
                nr.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            p_to_remove = last_p._p
            p_to_remove.getparent().remove(p_to_remove)
            
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Bullet List (* or -)
        if stripped.startswith(('* ', '- ', '• ')):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            # Parse bold tokens inside bullet
            tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', item_text)
            for tok in tokens:
                if not tok:
                    continue
                if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
                    run = p.add_run(tok[2:-2])
                    run.bold = True
                elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
                    run = p.add_run(tok[1:-1])
                    run.italic = True
                elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
                    run = p.add_run(tok[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                else:
                    p.add_run(tok)
            i += 1
            continue

        # Numbered List (1. , 2. )
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', item_text)
            for tok in tokens:
                if not tok:
                    continue
                if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
                    run = p.add_run(tok[2:-2])
                    run.bold = True
                elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
                    run = p.add_run(tok[1:-1])
                    run.italic = True
                elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
                    run = p.add_run(tok[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                else:
                    p.add_run(tok)
            i += 1
            continue

        # Math equation line ($$...$$)
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            eq_text = stripped[2:-2].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "FAFBFD")
            set_cell_margins(cell, top=80, bottom=80, left=160, right=160)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(eq_text)
            run.font.name = 'Cambria Math'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            run.italic = True
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Regular paragraph
        add_styled_paragraph(doc, stripped, space_after=6)
        i += 1

    # Cleanup remaining table/code
    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Successfully created Word Document: {output_path}")

if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    
    # 1. Primary Research Paper
    md_file = root / "research_paper" / "RESEARCH_PAPER.md"
    out_file1 = root / "research_paper" / "RESEARCH_PAPER.docx"
    out_file2 = root / "research paper" / "RESEARCH_PAPER.docx"
    try:
        create_research_docx(md_file, out_file1)
        create_research_docx(md_file, out_file2)
    except PermissionError:
        print("[Notice] RESEARCH_PAPER.docx is currently open in Word. Skipping overwrite.")

    # 2. Complete Journal Research Paper
    j_md_file = root / "research_paper" / "JOURNAL_RESEARCH_PAPER.md"
    j_out_file1 = root / "research_paper" / "JOURNAL_RESEARCH_PAPER.docx"
    j_out_file2 = root / "research paper" / "JOURNAL_RESEARCH_PAPER.docx"
    try:
        create_research_docx(j_md_file, j_out_file1)
        create_research_docx(j_md_file, j_out_file2)
    except PermissionError:
        print("[Notice] JOURNAL_RESEARCH_PAPER.docx is currently open in Word. Skipping overwrite.")
